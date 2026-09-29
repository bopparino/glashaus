"""Glashaus document renderer. Only validated structured content crosses here.

Reuses MIT-licensed Hermes docx block primitives, not its agent runtime.
Both formats come from the same document so page content stays consistent.
"""
import json
from pathlib import Path
import subprocess
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
sys.path.insert(0, '/opt/glashaus/vendor/hermes')
from docx_create import add_block

plan = json.loads(Path('/work/plan.json').read_text())
doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor.from_string('202320')
    style.font.size = Pt({'Title': 26, 'Heading 1': 16, 'Heading 2': 13}.get(name, 11))
    style.paragraph_format.space_after = Pt(8)
    if name.startswith('Heading'):
        style.paragraph_format.space_before = Pt(16)
        style.paragraph_format.keep_with_next = True
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.15
for section in doc.sections:
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
doc.core_properties.title = plan['title']
doc.core_properties.author = 'Glashaus'
doc.add_paragraph(plan['title'], 'Title')
doc.add_paragraph(plan['summary'])
for section in plan['sections']:
    add_block(doc, {'type': 'heading', 'text': section['heading'], 'level': 1})
    for paragraph in section['paragraphs']:
        add_block(doc, {'type': 'paragraph', 'text': paragraph})
    if section['items']:
        add_block(doc, {'type': 'bullet_list', 'items': section['items']})
if plan['caveats']:
    add_block(doc, {'type': 'heading', 'text': 'Notes and assumptions', 'level': 1})
    for note in plan['caveats']:
        add_block(doc, {'type': 'paragraph', 'text': note})
out = Path('/work/output')
out.mkdir(exist_ok=True)
doc.save(out / 'plan.docx')
subprocess.run(['libreoffice', '-env:UserInstallation=file:///tmp/lo-profile', '--headless', '--convert-to', 'pdf', '--outdir', str(out), str(out/'plan.docx')], check=True, timeout=60)
if not (out/'plan.pdf').is_file():
    raise RuntimeError('The PDF renderer did not produce a document')
from pypdf import PdfReader
if len(PdfReader(out/'plan.pdf').pages) > plan.get('maxPages', 8):
    raise RuntimeError('The plan exceeds the requested page limit. Retry with a shorter scope; no files have been published.')
